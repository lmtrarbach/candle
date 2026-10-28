"""Manuscript export: Markdown, HTML, plain text, ODT, PDF, DOCX and EPUB."""

import datetime
import html
import uuid
import zipfile
from dataclasses import dataclass

from PySide6.QtCore import QMarginsF
from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter, QTextDocument, QTextDocumentWriter

from . import richtext

SCENE_BREAK = "* * *"

FORMATS = {
    "Markdown (.md)": "md",
    "HTML (.html)": "html",
    "Plain text (.txt)": "txt",
    "OpenDocument (.odt)": "odt",
    "PDF (.pdf)": "pdf",
    "Word (.docx)": "docx",
    "EPUB e-book (.epub)": "epub",
}


@dataclass
class Options:
    chapter_titles: bool = True
    number_chapters: bool = True
    title_page: bool = True
    only_done: bool = False


@dataclass
class Chapter:
    title: str
    scenes: list   # list[list[Para]]


def collect(project, options):
    chapters = []
    for number, ch in enumerate(project.list("chapters"), start=1):
        scenes = []
        for sc in project.list("scenes", chapter_id=ch["id"]):
            if options.only_done and sc["status"] != "done":
                continue
            paras = richtext.non_empty(richtext.parse(sc["content"]))
            if paras:
                scenes.append(paras)
        if options.number_chapters:
            title = f"{number}. {ch['title']}" if ch["title"] else f"Chapter {number}"
        else:
            title = ch["title"]
        chapters.append(Chapter(title, scenes))
    return chapters


def export(project, path, fmt, options=None):
    options = options or Options()
    chapters = collect(project, options)
    meta = {
        "title": project.title,
        "author": project.get_meta("author", "") or "",
        "language": project.get_meta("language", "") or "en",
    }
    writer = {
        "md": write_markdown, "html": write_html, "txt": write_text, "odt": write_odt,
        "pdf": write_pdf, "docx": write_docx, "epub": write_epub,
    }[fmt]
    writer(path, meta, chapters, options)


# -- text formats ------------------------------------------------------------

def write_markdown(path, meta, chapters, options):
    out = []
    if options.title_page:
        out.append(f"# {meta['title']}")
        if meta["author"]:
            out.append(f"*{meta['author']}*")
    for ch in chapters:
        if options.chapter_titles and ch.title:
            out.append(f"## {ch.title}")
        out.append(f"\n\n{SCENE_BREAK}\n\n".join(richtext.to_markdown(s, 2) for s in ch.scenes))
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n\n".join(x for x in out if x) + "\n")


def write_text(path, meta, chapters, options):
    out = []
    if options.title_page:
        out.append(meta["title"].upper())
        if meta["author"]:
            out.append(meta["author"])
    for ch in chapters:
        if options.chapter_titles and ch.title:
            out.append(f"\n{ch.title}\n{'=' * len(ch.title)}")
        out.append(f"\n\n{SCENE_BREAK}\n\n".join(richtext.to_text(s) for s in ch.scenes))
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n\n".join(x for x in out if x) + "\n")


CSS = """
body { font-family: Georgia, 'Times New Roman', serif; line-height: 1.6;
       max-width: 40em; margin: 2em auto; padding: 0 1em; }
h1.book-title { text-align: center; margin-top: 30%; font-size: 2.4em; }
p.author { text-align: center; font-style: italic; font-size: 1.3em; }
h2 { text-align: center; margin-top: 3em; page-break-before: always; }
p { margin: 0; text-indent: 1.5em; text-align: justify; }
p.break { text-align: center; text-indent: 0; margin: 1em 0; }
"""


def _chapter_html(ch, options, heading_tag="h2"):
    parts = []
    if options.chapter_titles and ch.title:
        parts.append(f"<{heading_tag}>{html.escape(ch.title)}</{heading_tag}>")
    scenes = [richtext.to_html(s, 2) for s in ch.scenes]
    parts.append(f'\n<p class="break">{SCENE_BREAK}</p>\n'.join(scenes))
    return "\n".join(parts)


def _title_html(meta):
    out = f'<h1 class="book-title">{html.escape(meta["title"])}</h1>'
    if meta["author"]:
        out += f'\n<p class="author">{html.escape(meta["author"])}</p>'
    return out


def manuscript_html(meta, chapters, options):
    body = []
    if options.title_page:
        body.append(_title_html(meta))
    body.extend(_chapter_html(ch, options) for ch in chapters)
    return (
        "<!DOCTYPE html>\n<html><head><meta charset=\"utf-8\"/>"
        f"<title>{html.escape(meta['title'])}</title><style>{CSS}</style></head>\n"
        "<body>\n" + "\n".join(body) + "\n</body></html>\n"
    )


def write_html(path, meta, chapters, options):
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(manuscript_html(meta, chapters, options))


# -- Qt-rendered formats -----------------------------------------------------

def _qt_document(meta, chapters, options):
    doc = QTextDocument()
    doc.setDefaultStyleSheet(CSS)
    doc.setHtml(manuscript_html(meta, chapters, options))
    doc.setMetaInformation(QTextDocument.MetaInformation.DocumentTitle, meta["title"])
    return doc


def write_odt(path, meta, chapters, options):
    writer = QTextDocumentWriter(path, b"ODF")
    if not writer.write(_qt_document(meta, chapters, options)):
        raise OSError(f"Could not write {path}")


def write_pdf(path, meta, chapters, options):
    writer = QPdfWriter(path)
    writer.setTitle(meta["title"])
    writer.setCreator("Candle")
    writer.setPageLayout(QPageLayout(
        QPageSize(QPageSize.PageSizeId.A5), QPageLayout.Orientation.Portrait,
        QMarginsF(18, 20, 18, 20), QPageLayout.Unit.Millimeter,
    ))
    doc = _qt_document(meta, chapters, options)
    doc.print_(writer)


# -- DOCX --------------------------------------------------------------------

def write_docx(path, meta, chapters, options):
    try:
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
        from docx.shared import Pt
    except ImportError as exc:
        raise RuntimeError("DOCX export needs python-docx: pip install python-docx") from exc

    aligns = {
        "left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    document = docx.Document()
    document.core_properties.title = meta["title"]
    document.core_properties.author = meta["author"]
    normal = document.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(12)

    first = True
    if options.title_page:
        document.add_heading(meta["title"], 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
        if meta["author"]:
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(meta["author"]).italic = True
        first = False

    for ch in chapters:
        if not first:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        first = False
        if options.chapter_titles and ch.title:
            document.add_heading(ch.title, 1).alignment = WD_ALIGN_PARAGRAPH.CENTER
        for i, scene in enumerate(ch.scenes):
            if i:
                document.add_paragraph(SCENE_BREAK).alignment = WD_ALIGN_PARAGRAPH.CENTER
            for para in scene:
                if para.kind.startswith("h"):
                    p = document.add_heading(level=min(int(para.kind[1]) + 1, 9))
                elif para.kind == "ul":
                    p = document.add_paragraph(style="List Bullet")
                elif para.kind == "ol":
                    p = document.add_paragraph(style="List Number")
                else:
                    p = document.add_paragraph()
                    p.paragraph_format.first_line_indent = Pt(18)
                p.alignment = aligns[para.align]
                for run in para.runs:
                    r = p.add_run(run.text)
                    r.bold, r.italic, r.underline = run.bold, run.italic, run.underline
                    r.font.strike = run.strike
    document.save(path)


# -- EPUB 3 ------------------------------------------------------------------

def _xhtml(title, body, lang):
    return (
        '<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n'
        f'<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" '
        f'xml:lang="{lang}" lang="{lang}">\n<head><meta charset="utf-8"/>'
        f'<title>{html.escape(title)}</title>'
        '<link rel="stylesheet" type="text/css" href="style.css"/></head>\n'
        f"<body>\n{body}\n</body>\n</html>\n"
    )


def write_epub(path, meta, chapters, options):
    lang = html.escape(meta.get("language") or "en")
    book_id = f"urn:uuid:{uuid.uuid4()}"
    modified = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    pages = []  # (filename, title, xhtml)
    if options.title_page:
        pages.append(("title.xhtml", meta["title"], _xhtml(meta["title"], _title_html(meta), lang)))
    for i, ch in enumerate(chapters, start=1):
        title = ch.title or f"Chapter {i}"
        pages.append((f"chapter{i:03d}.xhtml", title, _xhtml(title, _chapter_html(ch, options, "h1"), lang)))

    nav_items = "\n".join(
        f'<li><a href="{name}">{html.escape(title)}</a></li>' for name, title, _ in pages
    )
    nav = _xhtml("Contents", f'<nav epub:type="toc" id="toc"><h1>Contents</h1><ol>\n{nav_items}\n</ol></nav>', lang)
    manifest = "\n".join(
        f'<item id="p{i}" href="{name}" media-type="application/xhtml+xml"/>'
        for i, (name, _, _) in enumerate(pages)
    )
    spine = "\n".join(f'<itemref idref="p{i}"/>' for i in range(len(pages)))
    opf = f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid" xml:lang="{lang}">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
<dc:identifier id="bookid">{book_id}</dc:identifier>
<dc:title>{html.escape(meta['title'])}</dc:title>
<dc:creator>{html.escape(meta['author'] or 'Unknown')}</dc:creator>
<dc:language>{lang}</dc:language>
<meta property="dcterms:modified">{modified}</meta>
</metadata>
<manifest>
<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>
<item id="css" href="style.css" media-type="text/css"/>
{manifest}
</manifest>
<spine>
{spine}
</spine>
</package>
"""
    container = """<?xml version="1.0" encoding="utf-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>
"""
    epub_css = CSS.replace("max-width: 40em; margin: 2em auto; padding: 0 1em;", "")
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("META-INF/container.xml", container, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/content.opf", opf, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/nav.xhtml", nav, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/style.css", epub_css, compress_type=zipfile.ZIP_DEFLATED)
        for name, _, body in pages:
            zf.writestr(f"OEBPS/{name}", body, compress_type=zipfile.ZIP_DEFLATED)
